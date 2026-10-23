import itertools

import openpyxl

DOSSIER = "sources/projections"
CENTRAL = "00_central"
PASSE = 1962
DEBUT = 2023
FIN = 2070
AGE_MAX = 120

HYPOTHESES = [
    ("fecondite", "Fécondité (enfants par femme)", [
        ("basse", "1,20", "12_fec_basse"),
        ("centrale", "1,45", CENTRAL),
        ("haute", "1,70", "11_fec_haute"),
        ]),
    ("mortalite", "Espérance de vie à la naissance en 2070 (femmes / hommes)", [
        ("basse", "86,5 / 83,7 ans", "14_ev_basse"),
        ("centrale", "89,5 / 86,7 ans", CENTRAL),
        ("haute", "92,5 / 89,7 ans", "13_ev_haute"),
        ]),
    ("migrations", "Solde migratoire annuel", [
        ("basse", "+70 000", "16_smi_bas"),
        ("centrale", "+150 000", CENTRAL),
        ("haute", "+230 000", "15_smi_haut"),
        ]),
    ]

ONGLETS = {
    "fecondite": [("hyp_fecondite", 2, 36)],
    "mortalite": [("hyp_mortaliteF", 3, 121), ("hyp_mortaliteH", 3, 121)],
    "migrations": [("hyp_soldemigF", 2, 121), ("hyp_soldemigH", 2, 121)],
    }

TRANCHES = [(18, 29), (30, 39), (40, 49), (50, 64), (65, AGE_MAX)]

MOTEURS = {
    "Enseignement": (3, 24),
    "Famille": (0, 19),
    }


def colonnes(entete):
    return {annee: entete.index(annee) for annee in range(PASSE, FIN + 1)}


def lire_ages(classeur, onglet, premiere, nombre):
    lignes = list(classeur[onglet].iter_rows(values_only=True))
    return {
        annee: [float(ligne[colonne] or 0) for ligne in lignes[premiere:premiere + nombre]]
        for annee, colonne in colonnes(lignes[1]).items()
        }


def lire_ligne(classeur, onglet, libelle):
    lignes = list(classeur[onglet].iter_rows(values_only=True))
    trouvee = next(ligne for ligne in lignes if str(ligne[0]).strip() == libelle)
    return {annee: float(trouvee[colonne]) for annee, colonne in colonnes(lignes[1]).items()}


def ouvrir(fichier):
    return openpyxl.load_workbook(f"{DOSSIER}/{fichier}.xlsx", read_only=True, data_only=True)


def etendre(effectifs):
    reste = effectifs[-1] / (AGE_MAX - 104)
    return effectifs[:-1] + [reste] * (AGE_MAX - 104)


def lire_commun():
    classeur = ouvrir(CENTRAL)
    naissances = lire_ligne(classeur, "naissance", "Total des naissances")
    filles = lire_ligne(classeur, "naissance", "dont naissances de filles")
    femmes = lire_ages(classeur, "populationF", 2, 106)
    hommes = lire_ages(classeur, "populationH", 2, 106)
    return {
        "observee": {
            annee: {"F": etendre(femmes[annee]), "H": etendre(hommes[annee])}
            for annee in range(PASSE, DEBUT + 1)
            },
        "ajustement": {
            "F": lire_ages(classeur, "ajustementF", 2, AGE_MAX + 1),
            "H": lire_ages(classeur, "ajustementH", 2, AGE_MAX + 1),
            },
        "part_filles": {annee: filles[annee] / naissances[annee] for annee in naissances},
        }


def lire_hypotheses():
    resultat = {}
    for composante, _, variantes in HYPOTHESES:
        for variante, _, fichier in variantes:
            classeur = ouvrir(fichier)
            tables = [lire_ages(classeur, onglet, premiere, nombre) for onglet, premiere, nombre in ONGLETS[composante]]
            resultat[composante, variante] = tables[0] if len(tables) == 1 else {"F": tables[0], "H": tables[1]}
    return resultat


def avancer(effectifs, fecondite, quotients, soldes, ajustements, part_filles):
    suivants = {}
    for sexe in ("F", "H"):
        suivants[sexe] = [0] + [
            effectifs[sexe][age - 1] * (1 - quotients[sexe][age] / 1e5) + soldes[sexe][age] + ajustements[sexe][age]
            for age in range(1, AGE_MAX + 1)
            ]
    meres = [(effectifs["F"][age - 1] + suivants["F"][age]) / 2 for age in range(15, 51)]
    naissances = sum(taux / 1e4 * n for taux, n in zip(fecondite, meres))
    for sexe, part in (("F", part_filles), ("H", 1 - part_filles)):
        nes = naissances * part * (1 - quotients[sexe][0] / 1e5)
        suivants[sexe][0] = nes + soldes[sexe][0] + ajustements[sexe][0]
    return suivants


def projeter(commun, fecondite, mortalite, migrations):
    population = dict(commun["observee"])
    for annee in range(DEBUT, FIN + 1):
        population[annee + 1] = avancer(
            population[annee],
            fecondite[annee],
            {sexe: mortalite[sexe][annee] for sexe in ("F", "H")},
            {sexe: migrations[sexe][annee] for sexe in ("F", "H")},
            {sexe: commun["ajustement"][sexe][annee] for sexe in ("F", "H")},
            commun["part_filles"][annee],
            )
    return population


def moyenne_annuelle(population, annee):
    return [
        (population[annee]["F"][age] + population[annee]["H"][age]
         + population[annee + 1]["F"][age] + population[annee + 1]["H"][age]) / 2
        for age in range(AGE_MAX + 1)
        ]


def effectif(ages, bornes):
    return sum(ages[bornes[0]:bornes[1] + 1])


def facteurs(population, annees):
    depart = moyenne_annuelle(population, DEBUT)
    resultat = []
    for annee in annees:
        ages = moyenne_annuelle(population, annee)
        resultat.append({
            "tranches": [effectif(ages, t) / effectif(depart, t) for t in TRANCHES],
            "moteurs": {libelle: effectif(ages, b) / effectif(depart, b) for libelle, b in MOTEURS.items()},
            "population": sum(ages) / sum(depart),
            "age65": effectif(ages, (65, AGE_MAX)) / sum(ages),
            })
    return resultat


def projeter_montants(profils, revenu, primaire, individus, libelles, annuels):
    lignes = []
    for profil, libelle in zip(profils, libelles):
        lignes.append([
            sum(profil) * f["moteurs"][libelle] if libelle in MOTEURS
            else sum(p * g for p, g in zip(profil, f["tranches"]))
            for f in annuels
            ])
    reste = [r - sum(colonne) for r, colonne in zip(revenu, zip(*profils))]
    return {
        "lignes": [[round(v, 2) for v in ligne] for ligne in lignes],
        "revenu": [
            round(sum(r * g for r, g in zip(reste, f["tranches"])) + sum(ligne[j] for ligne in lignes), 2)
            for j, f in enumerate(annuels)
            ],
        "primaire": [round(sum(p * g for p, g in zip(primaire, f["tranches"])), 2) for f in annuels],
        "individus": [round(sum(individus) * f["population"], 3) for f in annuels],
        "age65": [round(f["age65"], 4) for f in annuels],
        }


def scenarios(profils, revenu, primaire, individus, libelles, annees):
    commun = lire_commun()
    hypotheses = lire_hypotheses()
    variantes = [[variante for variante, _, _ in liste] for _, _, liste in HYPOTHESES]
    resultat = {}
    for fecondite, mortalite, migrations in itertools.product(*variantes):
        population = projeter(
            commun,
            hypotheses["fecondite", fecondite],
            hypotheses["mortalite", mortalite],
            hypotheses["migrations", migrations],
            )
        resultat[f"{fecondite}-{mortalite}-{migrations}"] = projeter_montants(
            profils, revenu, primaire, individus, libelles, facteurs(population, annees),
            )
    return resultat
